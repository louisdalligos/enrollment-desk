"""Deterministic roster parsing for CSV and Excel files.

Phase 1 maps headers with a synonym table. Headers it can't map are returned in
`unmapped_headers`; phase 2 hands only those to the AI column-mapping step, and
rows are always parsed in code, never by the model.
"""

from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from pathlib import Path

from .models import Flag, MemberRecord, ParseResult, Relationship, Severity
from .reference import plan_by_code_or_name

HEADER_SYNONYMS: dict[str, list[str]] = {
    "employee_id": ["employee id", "emp id", "emp no", "employee no", "employee number", "id no", "staff id", "emp number"],
    "first_name": ["first name", "firstname", "given name", "fname"],
    "last_name": ["last name", "lastname", "surname", "family name", "lname"],
    "full_name": ["name", "full name", "employee name", "member name", "name last first"],
    "birth_date": ["birth date", "birthdate", "date of birth", "dob", "bday", "birthday"],
    "sex": ["sex", "gender"],
    "relationship": ["relationship", "relation", "member type", "rel"],
    "principal_employee_id": ["principal id", "principal employee id", "principal emp id", "dependent of", "principal emp no"],
    "hire_date": ["hire date", "date hired", "start date", "date of hire", "employment date"],
    "plan_code": ["plan", "plan code", "plan type", "benefit plan"],
}
_LOOKUP = {syn: field for field, syns in HEADER_SYNONYMS.items() for syn in syns}

DATE_FIELDS = ("birth_date", "hire_date")
DATE_FORMATS = ("%Y-%m-%d", "%m/%d/%Y", "%Y/%m/%d", "%d-%b-%Y", "%B %d, %Y", "%b %d, %Y", "%d %B %Y", "%m-%d-%Y")

SEX_VALUES = {"m": "M", "male": "M", "lalaki": "M", "f": "F", "female": "F", "babae": "F"}
REL_VALUES = {
    "principal": Relationship.PRINCIPAL, "employee": Relationship.PRINCIPAL, "self": Relationship.PRINCIPAL,
    "emp": Relationship.PRINCIPAL, "spouse": Relationship.SPOUSE, "wife": Relationship.SPOUSE,
    "husband": Relationship.SPOUSE, "asawa": Relationship.SPOUSE, "child": Relationship.CHILD,
    "son": Relationship.CHILD, "daughter": Relationship.CHILD, "anak": Relationship.CHILD,
    "parent": Relationship.PARENT, "mother": Relationship.PARENT, "father": Relationship.PARENT,
}


def normalize_header(h: str) -> str:
    h = str(h).strip().lower().replace("#", " no ")
    h = re.sub(r"[_\-./:()]", " ", h)
    h = re.sub(r"[^a-z0-9 ]", "", h)
    return re.sub(r"\s+", " ", h).strip()


def map_header(h: str) -> str | None:
    return _LOOKUP.get(normalize_header(h))


def parse_date(value) -> date | None:
    """Parse a cell into a date. Raises ValueError if the text is not a recognised date."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = re.sub(r"\s+", " ", str(value).strip())
    if text.endswith(" 00:00:00"):
        text = text[: -len(" 00:00:00")]
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"unrecognised date: {text!r}")


def clean_name(value) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    if not text:
        return None
    if text.isupper() or text.islower():
        text = " ".join(w.capitalize() for w in text.split(" "))
    return text


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def read_rows(path: str | Path, content: bytes | None = None) -> list[list]:
    """Read a CSV or XLSX file into a list of rows (cells keep native types for Excel)."""
    path = Path(path)
    data = content if content is not None else path.read_bytes()
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        try:
            ws = wb.worksheets[0]
            if (ws.max_row and ws.max_row > 1015) or (ws.max_column and ws.max_column > 100):
                raise ValueError("roster exceeds 1000 data rows or 100 columns")
            rows = []
            for row in ws.iter_rows(values_only=True):
                if len(rows) >= 1015 or len(row) > 100:
                    raise ValueError("roster exceeds 1000 data rows or 100 columns")
                rows.append(list(row))
            return rows
        finally:
            wb.close()
    text = data.decode("utf-8-sig")
    try:
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = []
    for row in csv.reader(io.StringIO(text), dialect):
        if len(rows) >= 1015 or len(row) > 100:
            raise ValueError("roster exceeds 1000 data rows or 100 columns")
        rows.append(row)
    return rows


def _find_header(rows: list[list]) -> int:
    """Index of the first row with at least three cells that map to known fields (skips title rows)."""
    for i, row in enumerate(rows[:15]):
        if sum(1 for c in row if c is not None and map_header(str(c))) >= 3:
            return i
    raise ValueError("no header row found: fewer than three recognisable column names in the first 15 rows")


def parse_roster(path: str | Path, content: bytes | None = None, *, header_row: int | None = None, column_map: dict | None = None) -> ParseResult:
    rows = read_rows(path, content)
    h_idx = _find_header(rows) if header_row is None else header_row
    if header_row is not None:
        from .column_mapping import headers_for, validate_mapping
        headers = headers_for(rows, header_row)
        if column_map is None:
            raise ValueError("Explicit header rows require a reviewed column mapping")
        validate_mapping(headers, column_map)
    else:
        if column_map is not None:
            raise ValueError("A reviewed mapping requires an explicit header row")
        headers = [_cell_text(c) for c in rows[h_idx]]
    header_map: dict[str, str] = {}
    unmapped: list[str] = []
    col_field: dict[int, str] = {}
    for i, h in enumerate(headers):
        if not h:
            continue
        field = column_map.get(str(i)) if column_map is not None else map_header(h)
        if field and field not in col_field.values():
            col_field[i] = field
            header_map[h] = field
        else:
            unmapped.append(h)

    members: list[MemberRecord] = []
    flags: list[Flag] = []
    data_row = 0
    for raw_row in rows[h_idx + 1 :]:
        if all(_cell_text(c) == "" for c in raw_row):
            continue
        data_row += 1
        raw = {headers[i]: _cell_text(c) for i, c in enumerate(raw_row) if i < len(headers) and headers[i]}
        values: dict[str, object] = {}
        for i, field in col_field.items():
            cell = raw_row[i] if i < len(raw_row) else None
            if field in DATE_FIELDS:
                try:
                    values[field] = parse_date(cell)
                except ValueError:
                    flags.append(Flag(code="unparseable_date", row=data_row, field=field,
                                      message=f"Could not read {field.replace('_', ' ')} '{_cell_text(cell)}'"))
            else:
                values[field] = _cell_text(cell) or None

        full = values.pop("full_name", None)
        if full and not (values.get("first_name") or values.get("last_name")):
            if "," in str(full):
                last, first = str(full).split(",", 1)
            else:
                parts = str(full).rsplit(" ", 1)
                first, last = (parts[0], parts[1]) if len(parts) == 2 else (parts[0], "")
            values["first_name"], values["last_name"] = first, last

        sex = values.get("sex")
        if sex is not None:
            mapped = SEX_VALUES.get(str(sex).strip().lower())
            if mapped is None:
                flags.append(Flag(code="invalid_value", row=data_row, field="sex", message=f"Unknown sex '{sex}'"))
            values["sex"] = mapped

        rel = values.get("relationship")
        if rel is not None:
            mapped_rel = REL_VALUES.get(str(rel).strip().lower())
            if mapped_rel is None:
                flags.append(Flag(code="invalid_value", row=data_row, field="relationship",
                                  message=f"Unknown relationship '{rel}'"))
            values["relationship"] = mapped_rel

        members.append(
            MemberRecord(
                row=data_row,
                employee_id=(values.get("employee_id") or None),
                first_name=clean_name(values.get("first_name")),
                last_name=clean_name(values.get("last_name")),
                birth_date=values.get("birth_date"),
                sex=values.get("sex"),
                relationship=values.get("relationship"),
                principal_employee_id=(values.get("principal_employee_id") or None),
                hire_date=values.get("hire_date"),
                plan_code=plan_by_code_or_name(values.get("plan_code")),
                raw=raw,
            )
        )
    for f in flags:
        f.severity = Severity.BLOCK
    return ParseResult(members=members, flags=flags, header_map=header_map, unmapped_headers=unmapped)
